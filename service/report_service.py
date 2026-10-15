import time
import os
import json
import re
# ================= 报告生成重依赖（matplotlib / python-docx）=================
# 海外 LITE 精简环境缺失时优雅降级，不影响系统启动
try:
    import matplotlib.pyplot as plt

    plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei", "Arial Unicode MS"]
    plt.rcParams["axes.unicode_minus"] = False
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    _REPORT_DOC_AVAILABLE = True
except Exception as _report_dep_err:
    plt = None
    Document = None
    Pt = RGBColor = Cm = Inches = None
    WD_ALIGN_PARAGRAPH = None
    qn = None
    _REPORT_DOC_AVAILABLE = False
    print(f"[LITE] 报告生成组件不可用（{_report_dep_err}），Word/图表报告功能将降级")

from service.xinghuo_llm import xinghuo_llm
from service.rag_service import get_knowledge, rag_db


# ==================== 敏感词列表 ====================
SENSITIVE_WORDS = [
    "亚文化", "解构", "话语体系", "群体性事件", "极端主义",
    "政治", "宗教", "民粹", "异见", "敏感", "颠覆", "推翻",
    "阶级斗争", "颜色革命", "人权", "民主运动"
]


def extract_fact_sheet(use_data):
    facts = {}
    overview = use_data.get("overview_stats", {})
    facts["total_articles"] = overview.get("total_articles", 0)
    facts["total_comments"] = overview.get("total_comments", 0)
    facts["real_latest_date"] = overview.get("real_latest_date", "未知")
    facts["real_earliest_date"] = overview.get("real_earliest_date", "未知")

    sent = use_data.get("sentiment_trend", {})
    facts["pos_ratio"] = sent.get("pos_ratio", "0%")
    facts["neu_ratio"] = sent.get("neu_ratio", "0%")
    facts["neg_ratio"] = sent.get("neg_ratio", "0%")
    facts["sent_total"] = sent.get("total", 0)

    hot_topics = use_data.get("hot_topics", [])[:5]
    facts["top_keywords"] = [t["keyword"] for t in hot_topics]
    facts["top_heats"] = [t["heat"] for t in hot_topics]
    facts["top_sentiments"] = [t.get("sentiment", "中性") for t in hot_topics]

    trend = use_data.get("daily_trend", [])
    if trend:
        counts = [t["count"] for t in trend]
        max_idx = counts.index(max(counts))
        min_idx = counts.index(min(counts))
        facts["peak_date"] = trend[max_idx]["date"][-5:]
        facts["peak_count"] = max(counts)
        facts["valley_date"] = trend[min_idx]["date"][-5:]
        facts["valley_count"] = min(counts)
    else:
        facts["peak_date"] = facts["valley_date"] = "无数据"
        facts["peak_count"] = facts["valley_count"] = 0

    return facts


def format_data_for_llm(data):
    if isinstance(data, dict):
        return "\n".join([f"- {k}: {v}" for k, v in data.items()])
    elif isinstance(data, list):
        clean_items = []
        for i, item in enumerate(data[:20]):
            if isinstance(item, dict):
                props = ", ".join([f"{k}:{v}" for k, v in item.items()])
                clean_items.append(f"{i + 1}. {props}")
            else:
                clean_items.append(f"{i + 1}. {str(item)}")
        return "\n".join(clean_items)
    return str(data)


def generate_charts(use_data, report_dir="reports"):
    """生成图表。只使用真实数据；数据缺失时跳过该图，绝不编造示例数据。"""
    if plt is None:
        print("[LITE] matplotlib 不可用，跳过图表生成")
        return {}
    os.makedirs(report_dir, exist_ok=True)
    chart_paths = {}

    trend = use_data.get("daily_trend", [])
    sent = use_data.get("sentiment_trend", {})
    hot_topics = use_data.get("hot_topics", [])
    comments = use_data.get("user_comments", [])
    trend_is_real = bool(trend)
    sent_is_real = bool(sent.get("total", 0))
    hot_is_real = bool(hot_topics)
    comments_is_real = bool(comments)

    # ---------- 1. 发文趋势（粒度自适应：月 / 日） ----------
    trend_grain = use_data.get("trend_grain", "day")
    _grain_cn = "月度" if trend_grain == "month" else "每日"
    if trend_is_real:
        try:
            # 月粒度 key = 'YYYY-MM'（7 字符）；日粒度 key = 'YYYY-MM-DD'，只显示后 5 位
            dates = [t["date"] if trend_grain == "month" else t["date"][-5:] for t in trend]
            counts = [t["count"] for t in trend]
            plt.figure(figsize=(9, 4))
            plt.plot(dates, counts, marker='o', color='#2980B9', linewidth=2)
            plt.title(f"{_grain_cn}发文量趋势", fontsize=14)
            plt.ylabel("发文量"); plt.grid(True, linestyle='--', alpha=0.5)
            plt.xticks(rotation=30, ha='right'); plt.tight_layout()
            path = os.path.join(report_dir, "chart_trend.png")
            plt.savefig(path, dpi=150); plt.close()
            chart_paths["trend"] = path
        except Exception as e:
            print(f"生成趋势图失败: {e}")
    else:
        print("[WARN] 无趋势数据，跳过趋势图（不编造示例数据）")

    # ---------- 2. 情感饼图 ----------
    if sent_is_real:
        try:
            sizes = [sent.get("positive", 0), sent.get("neutral", 0), sent.get("negative", 0)]
            if sum(sizes) > 0:
                labels = ["正面", "中性", "负面"]
                colors = ["#27AE60", "#F39C12", "#E74C3C"]
                plt.figure(figsize=(6, 5))
                plt.pie(sizes, labels=labels, colors=colors, autopct='%1.1f%%',
                        startangle=90, textprops={'fontsize': 12})
                plt.title("舆论情感占比分布", fontsize=14); plt.tight_layout()
                path = os.path.join(report_dir, "chart_sentiment.png")
                plt.savefig(path, dpi=150); plt.close()
                chart_paths["sentiment"] = path
        except Exception as e:
            print(f"生成情感图失败: {e}")
    else:
        print("[WARN] 无情感数据，跳过情感图")

    # ---------- 3. 热词柱状图 ----------
    if hot_is_real:
        try:
            top10 = hot_topics[:10]
            words = [t["keyword"] for t in top10]
            heats = [t["heat"] for t in top10]
            sentiments = [t.get("sentiment", "中性") for t in top10]
            color_map = {"正面": "#27AE60", "中性": "#F39C12", "负面": "#E74C3C"}
            bar_colors = [color_map.get(s, "#F39C12") for s in sentiments]
            plt.figure(figsize=(9, 4))
            plt.bar(words, heats, color=bar_colors)
            plt.title("TOP10热词热度与情感倾向", fontsize=14); plt.ylabel("热度值")
            plt.xticks(rotation=30, ha='right'); plt.grid(axis='y', linestyle='--', alpha=0.5); plt.tight_layout()
            path = os.path.join(report_dir, "chart_hotword_sentiment.png")
            plt.savefig(path, dpi=150); plt.close()
            chart_paths["hotword_sentiment"] = path
        except Exception as e:
            print(f"生成热词情感图失败: {e}")
    else:
        print("[WARN] 无热词数据，跳过热词图")

    # ---------- 4. 评论散点图 ----------
    if comments_is_real:
        try:
            likes = [c.get("likes", 0) for c in comments]
            lengths = [len(c.get("content", "")) for c in comments]
            plt.figure(figsize=(7, 4))
            plt.scatter(lengths, likes, alpha=0.6, color='#8E44AD', edgecolors='white')
            plt.title("评论互动质量分析", fontsize=14)
            plt.xlabel("评论内容长度（字符）"); plt.ylabel("点赞数")
            plt.grid(True, linestyle='--', alpha=0.5); plt.tight_layout()
            path = os.path.join(report_dir, "chart_comment_scatter.png")
            plt.savefig(path, dpi=150); plt.close()
            chart_paths["comment_scatter"] = path
        except Exception as e:
            print(f"生成评论散点图失败: {e}")
    else:
        print("[WARN] 无评论数据，跳过评论图")

    # ---------- 5. TOP10条形图 ----------
    if hot_is_real:
        try:
            top10 = list(reversed(hot_topics[:10]))
            words = [t["keyword"] for t in top10]
            heats = [t["heat"] for t in top10]
            plt.figure(figsize=(8, 4))
            plt.barh(words, heats, color='#2980B9')
            plt.title("热门议题TOP10排行", fontsize=14); plt.xlabel("热度值")
            plt.grid(axis='x', linestyle='--', alpha=0.5); plt.tight_layout()
            path = os.path.join(report_dir, "chart_top10_bar.png")
            plt.savefig(path, dpi=150); plt.close()
            chart_paths["top10_bar"] = path
        except Exception as e:
            print(f"生成TOP10条形图失败: {e}")

    print("[STAT] 已生成的图表：", list(chart_paths.keys()))
    return chart_paths

def clean_llm_output(text, chapter_title=""):
    if not text:
        return ""
    text = re.sub(r'<thought>.*?</thought>', '', text, flags=re.DOTALL)
    text = re.sub(r'(?m)^#+\s*', '', text)
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'(?m)^\s*[-*]\s+', '', text)

    lines = []
    for line in text.split('\n'):
        line = line.strip()
        num_match = re.match(r'^(\d+)[.)、\s]+(.*)$', line)
        if num_match:
            num = num_match.group(1)
            content = num_match.group(2).strip()
            line = f"{num}. {content}"
        lines.append(line)
    text = '\n'.join(lines)

    text = (text
            .replace("达到XX", "达到具体数值")
            .replace("高达XX", "高达具体数值")
            .replace("占比XX%", "占比具体百分比")
            .replace("XX篇", "具体篇数"))
    text = re.sub(r'(?<![a-zA-Z])XX(?![a-zA-Z])', '具体数值', text)

    if chapter_title:
        clean_title = re.sub(r'^[一二三四五六七八九十]+、\s*', '', chapter_title).strip()
        new_lines = []
        for line in text.split('\n'):
            stripped = line.strip()
            is_duplicate = (
                stripped == chapter_title or
                stripped == clean_title or
                (len(clean_title) > 3 and clean_title in stripped and len(stripped) < len(clean_title) + 10)
            )
            if not is_duplicate:
                new_lines.append(line)
        text = '\n'.join(new_lines)

    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def _set_run_font(run, cn_font='SimSun', en_font='Times New Roman', size=None, color=None):
    if not _REPORT_DOC_AVAILABLE:
        return
    run.font.name = en_font
    if size: run.font.size = size
    if color: run.font.color.rgb = color
    rPr = run.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), cn_font)


def deep_sanitize(text):
    """对文本进行敏感词替换 + 调用 xinghuo_llm._sanitize_text 双重清洗"""
    if not text:
        return text
    for word in SENSITIVE_WORDS:
        text = text.replace(word, "相关研究")
    if hasattr(xinghuo_llm, '_sanitize_text'):
        text = xinghuo_llm._sanitize_text(text, deep_mode=True)
    return text


def generate_full_report(use_data):
    if Document is None:
        print("[LITE] python-docx 不可用，返回降级提示")
        return "报告生成功能在当前部署环境不可用（海外演示版已精简组件，请在本机完整版生成报告）。"
    print("[START] 开始生成舆情分析报告...")
    start_time = time.time()

    fact_sheet = extract_fact_sheet(use_data)
    print(f"[OK] 事实清单提取完成: {json.dumps(fact_sheet, ensure_ascii=False)[:200]}...")

    print("[STAT] 正在生成数据可视化图表...")
    chart_paths = generate_charts(use_data)

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.space_before = Pt(0)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), 'SimSun')

    _basin = use_data.get('basin') if isinstance(use_data, dict) else None
    title = doc.add_heading((f'{_basin}流域舆情分析综合报告' if _basin else '网络舆情分析综合报告'), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    if title.runs:
        _set_run_font(title.runs[0], cn_font='SimHei', size=Pt(22), color=RGBColor(0, 0, 0))

    p_ai = doc.add_paragraph()
    p_ai.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ai.paragraph_format.space_after = Pt(24)
    run_ai = p_ai.add_run(
        '⚠️ [AI生成标识] 本报告由新闻传播学垂类大模型结合本地数据库与学科知识库(RAG)自动生成，仅供教学科研参考，不构成任何决策建议。')
    _set_run_font(run_ai, size=Pt(10), color=RGBColor(0x80, 0x80, 0x80))

    # 流域报告：注明统计口径（避免"总数=全议题、热词/情感=聚焦样本"被误读为矛盾）
    if use_data.get("basin"):
        p_scope = doc.add_paragraph()
        p_scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_scope.paragraph_format.space_after = Pt(18)
        _tnote = use_data.get("trend_note") or ""
        run_scope = p_scope.add_run(
            f'口径说明：本报告聚焦「{use_data["basin"]}流域」。入库总数 / 评论数为该流域全议题数据库统计；'
            f'热词、情感分布与趋势曲线基于「生态治理 + 文旅 + 医患矛盾」聚焦样本聚合，两者口径不同，请勿直接相除。'
            + (f'{_tnote}。' if _tnote else ''))
        _set_run_font(run_scope, size=Pt(9), color=RGBColor(0x80, 0x80, 0x80))

    safe_data = {}
    for key, value in use_data.items():
        safe_data[key] = format_data_for_llm(value)

    chapters = [
        {"title": "一、舆情整体态势与生命周期研判", "data_keys": ["overview_stats", "daily_trend"],
         "rag_query": "网络舆情生命周期理论 发酵期 爆发期", "chart": "trend"},
        {"title": "二、热门议题聚焦与传播层级分析", "data_keys": ["hot_topics", "keywords"],
         "rag_query": "议程设置理论 框架理论 社交媒体话题演化", "chart": "top10_bar"},
        {"title": "三、核心议题与情感倾向交叉分析", "data_keys": ["hot_topics", "sentiment_trend"],
         "rag_query": "情感传播机制 舆论极化 负面情绪动员", "chart": "hotword_sentiment"},
        {"title": "四、关键节点用户互动与意见领袖识别", "data_keys": ["user_comments", "overview_stats"],
         "rag_query": "意见领袖 两级传播 网络群体极化现象", "chart": "comment_scatter"},
        {"title": "五、舆论情感演化规律与风险预警", "data_keys": ["sentiment_trend", "daily_trend"],
         "rag_query": "网络舆情情感演化规律 沉默螺旋 风险感知", "chart": "sentiment"},
        {"title": "六、用户评论语言特征分析", "data_keys": ["user_comments", "keywords"],
         "rag_query": "网络语言特征 受众心理 使用与满足 社交媒体评论", "chart": None},
        {"title": "七、舆情综合研判与治理优化建议",
         "data_keys": ["overview_stats", "sentiment_trend", "hot_topics"],
         "rag_query": "企业舆情应对策略 危机公关 媒介化治理", "chart": None}
    ]

    total_chapters = len(chapters)
    for idx, chapter in enumerate(chapters):
        print(f"[{idx + 1}/{total_chapters}] 正在生成：{chapter['title']} ...")

        context_text = "暂无相关理论补充"
        if rag_db:
            try:
                raw_context = get_knowledge(chapter['rag_query'], rag_db, top_k=5)
                context_text = deep_sanitize(raw_context)
            except Exception as e:
                print(f"    [WARN] RAG 检索失败: {e}")

        chapter_data_list = []
        for key in chapter.get('data_keys', []):
            if key in safe_data:
                chapter_data_list.append(f"【{key}数据】:\n{safe_data[key]}")
        chapter_data = "\n\n".join(chapter_data_list) if chapter_data_list else '暂无数据'

        # 🌟 第六章特殊处理：过滤无意义的0值字段，避免大模型罗列
        if idx == 5:
            clean_fact = {}
            for k, v in fact_sheet.items():
                if isinstance(v, (int, float)) and v == 0:
                    continue
                if isinstance(v, str) and (v == "0%" or v == "0" or v == "无数据"):
                    continue
                if isinstance(v, list) and len(v) == 0:
                    continue
                clean_fact[k] = v
            # 确保关键字段存在
            for key in ["total_articles", "total_comments", "real_earliest_date", "real_latest_date"]:
                if key not in clean_fact and key in fact_sheet:
                    clean_fact[key] = fact_sheet[key]
            current_fact = clean_fact
        else:
            current_fact = fact_sheet

        final_prompt = f"""你是一位顶尖高校的【新闻传播学教授】兼【资深舆情分析专家】。
请根据【事实清单】、【参考数据】和【理论知识】，撰写报告章节：《{chapter['title']}》。

=== 【事实清单 - 数据基准】 ===
以下为系统已核实的数据；标注为“未知/无数据”的维度，请在正文中如实说明“该维度暂无完整数据”，严禁编造任何数值：
{json.dumps(current_fact, ensure_ascii=False, indent=2)}

=== 【理论知识】 ===
{context_text}

=== 【参考数据明细】 ===
{chapter_data}

=== 【严格写作要求】 ===
1. 【纯段落叙述】：必须写成连贯的学术分析长文，绝对禁止列表、编号、分点。
2. 【禁止数据堆砌】：不要用“XX：0%”这样的句式罗列数据，如果某项数据为0或未知，请直接用“该维度暂无显著特征/暂无完整数据”一笔带过。
3. 【三段式论证】：每个分析段落必须包含：①引用事实清单中的具体数值（选择有代表性的）；②用传播学理论解释该数据的成因；③给出明确的舆情判断。
4. 【时间锚点】：必须使用事实清单中的 real_earliest_date 和 real_latest_date；若为“未知”，则用“本期数据”等中性表述，禁止“近期”“最近”等模糊词或编造日期。
5. 【学术规范】：使用专业术语，关键论断后括号标注理论来源。
6. 【诚实原则】：任何数据必须来自事实清单，没有的数据一律不写、不猜、不补。\n7. 字数控制在 600-900 字。直接输出纯文本正文，不要标题，不要 <thought> 标签。"""

        # 🌟 对最终 prompt 再次清洗
        final_prompt = deep_sanitize(final_prompt)

        content = ""
        try:
            response = xinghuo_llm.client.chat.completions.create(
                model=xinghuo_llm.model_name,
                messages=[{"role": "user", "content": final_prompt}],
                temperature=0.3,
                max_tokens=2048
            )
            if response.choices and response.choices[0].message.content:
                content = response.choices[0].message.content.strip()
        except Exception as e:
            if "10013" in str(e) or "SuspectedResult" in str(e):
                print("    [WARN] 触发审核，尝试使用极简提示词重试...")
                safe_title = deep_sanitize(chapter['title'])
                simple_prompt = f"你是一位舆情分析师。请根据已有数据，撰写关于《{safe_title}》的简短段落，只讨论用户评论的语言特征和互动情况，不涉及任何敏感话题。字数400-500字，学术严谨。"
                simple_prompt = deep_sanitize(simple_prompt)
                try:
                    response = xinghuo_llm.client.chat.completions.create(
                        model=xinghuo_llm.model_name,
                        messages=[{"role": "user", "content": simple_prompt}],
                        temperature=0.3,
                        max_tokens=1024
                    )
                    content = response.choices[0].message.content.strip()
                except Exception as retry_err:
                    content = f"（该章节因合规限制未能生成，请检查数据或稍后重试）"
            else:
                content = f"（生成失败：{e}）"

        if not content:
            content = "（大模型返回内容为空，请重试）"

        clean_content = clean_llm_output(content, chapter_title=chapter['title'])
        if not clean_content:
            clean_content = "（大模型返回内容为空，请重试）"

        heading = doc.add_heading(chapter['title'], level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading.paragraph_format.space_before = Pt(18)
        heading.paragraph_format.space_after = Pt(12)
        heading.paragraph_format.first_line_indent = Cm(0)
        for run in heading.runs:
            _set_run_font(run, cn_font='SimHei', size=Pt(16), color=RGBColor(0, 0, 0))

        paragraphs = [p.strip() for p in clean_content.split('\n') if p.strip()]
        for para_text in paragraphs:
            p = doc.add_paragraph(para_text)
            p.style = doc.styles['Normal']
            p.paragraph_format.first_line_indent = Cm(0.74)
            p.paragraph_format.line_spacing = 1.5

        chart_key = chapter.get("chart")
        if chart_key and chart_key in chart_paths:
            try:
                img_para = doc.add_paragraph()
                img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_para.paragraph_format.space_before = Pt(12)
                img_para.paragraph_format.space_after = Pt(6)
                img_para.paragraph_format.first_line_indent = Cm(0)
                run_img = img_para.add_run()
                run_img.add_picture(chart_paths[chart_key], width=Inches(5.5))

                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.space_after = Pt(12)
                caption.paragraph_format.first_line_indent = Cm(0)
                cap_run = caption.add_run(f"图 {idx + 1}  {chapter['title']}数据可视化")
                _set_run_font(cap_run, size=Pt(9), color=RGBColor(0x66, 0x66, 0x66))
            except Exception as e:
                print(f"    [WARN] 插入图表失败: {e}")

        print(f"    [OK] {chapter['title']}")

    timestamp = time.strftime("%Y-%m-%d_%H-%M-%S")
    filename = f"reports/{(_basin + '流域') if _basin else ''}舆情分析报告_{timestamp}.docx"
    os.makedirs("reports", exist_ok=True)
    doc.save(filename)

    elapsed_time = time.time() - start_time
    print(f"\n[DONE] 报告生成完毕！总耗时: {elapsed_time:.1f} 秒")
    print(f"[FOLDER] 文件已保存至: {filename}")
    return filename