# -*- coding: utf-8 -*-
"""短视频智能分析 —— 专题简报（Word）生成器。

只依赖 python-docx（已装），不引入 matplotlib 等重依赖，**不改** service/report_service.py。
产物写入 static/reports/，复用既有的 /page/download_report 路由下载。

用法：
    python service/video/insight/report_docx.py --basin yz        # 长江
    python service/video/insight/report_docx.py --basin hh --top 15
"""
from __future__ import annotations

import argparse
import io
import json
import os
import time

_HERE = os.path.dirname(os.path.abspath(__file__))
_PRJ = os.path.dirname(os.path.dirname(os.path.dirname(_HERE)))
TOPIC = os.path.join(_PRJ, 'dashboard_gen', 'video_topic_data.json')
REPORTS = os.path.join(_PRJ, 'static', 'reports')

PFX2BASIN = {'yz': '长江', 'hh': '黄河'}


def _font(run, size=10.5, bold=False):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')


def _p(doc, text='', size=10.5, bold=False, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if text:
        _font(p.add_run(text), size, bold)
    return p


def _tbl(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(header):
        _font(t.rows[0].cells[i].paragraphs[0].add_run(str(h)), 10, True)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            _font(cells[i].paragraphs[0].add_run('' if v is None else str(v)), 9.5)
    return t


def build(basin, top=10, full_detail=True):
    """生成简报，返回写入的绝对路径。"""
    from docx import Document
    data = json.load(io.open(TOPIC, encoding='utf-8'))
    node = (data.get('topics') or {}).get(basin)
    if not node or not node.get('cards'):
        raise SystemExit('该专题暂无数据：%s' % basin)
    cards = node['cards']
    agg = node.get('agg') or {}
    meta = data.get('meta') or {}

    doc = Document()
    _p(doc, '%s · 短视频智能分析简报' % basin, 18, True)
    _p(doc, '生成时间：%s' % time.strftime('%Y-%m-%d %H:%M'),
       9.5).runs[0].font.color.rgb = None
    _p(doc, '数据口径：%s' % (meta.get('note') or ''), 9.5)

    # 一、样本概览
    _p(doc, '一、样本概览', 13, True)
    _tbl(doc, ['指标', '值'], [
        ['样本条数', node.get('n')],
        ['议题分布', '、'.join('%s %d' % (a, b) for a, b in (agg.get('topic') or []))],
        ['画幅构成', '、'.join('%s %d' % (a, b) for a, b in (agg.get('orient') or []))],
        ['画面主色调', '、'.join('%s %d' % (a, b) for a, b in (agg.get('hue') or [])[:8])],
        ['已生成推荐结论', '%d 条' % sum(1 for c in cards if (c.get('top') or {}).get('claim'))],
    ])

    # 二、推荐结论 Top N
    _p(doc, '二、推荐结论 Top %d（按推荐度排序）' % top, 13, True)
    ranked = sorted([c for c in cards if (c.get('top') or {}).get('recommend') is not None],
                    key=lambda c: -c['top']['recommend'])[:top]
    rows = []
    for i, c in enumerate(ranked, 1):
        tp = c['top']
        rows.append([i, '%.3f' % tp['recommend'], c.get('topic'), (c.get('title') or '')[:26],
                     tp.get('claim') or '', '、'.join(tp.get('evidence') or [])])
    _tbl(doc, ['#', '推荐度', '议题', '标题', '推荐结论', '证据'], rows)

    # 三、分议题要点
    _p(doc, '三、分议题要点', 13, True)
    for topic in ('生态治理', '文旅'):
        sub = [c for c in ranked if c.get('topic') == topic][:3]
        if not sub:
            continue
        _p(doc, '【%s】' % topic, 11, True)
        for c in sub:
            tp = c['top']
            _p(doc, '· （推荐度 %.3f）%s' % (tp['recommend'], tp.get('claim') or ''), 10)
            if c.get('advice'):
                for a in c['advice'][:2]:
                    _p(doc, '    → 建议：%s' % a, 9.5)

    # 四、逐条明细
    if full_detail:
        _p(doc, '四、逐条明细（全部 %d 条）' % len(cards), 13, True)
        rows = []
        for i, c in enumerate(cards, 1):
            v = c.get('vlm') or {}
            th = c.get('theme') or {}
            sh = c.get('shot') or {}
            tp = c.get('top') or {}
            rows.append([
                i, (c.get('title') or '')[:24], c.get('topic'),
                th.get('标签') or '', (v.get('主体') or '')[:34],
                '%s/%s/%s' % (sh.get('画幅') or '-', sh.get('主色调') or '-', sh.get('时长') or '-'),
                ('%.3f' % tp['recommend']) if tp.get('recommend') is not None else '-',
                (tp.get('claim') or '')[:60],
            ])
        _tbl(doc, ['#', '标题', '议题', '主题标签', '画面主体', '画幅/色调/时长', '推荐度', '推荐结论'],
             rows)

    # 五、口径与声明
    _p(doc, '五、口径与声明', 13, True)
    for s in [
        '样本为 L2 关键帧层抽样（8% 且非随机，按高互动 / 情感对立 / 高转发优先）并仅覆盖抖音平台，'
        '单条结论不构成总体规律。',
        '画面读数由视觉模型对关键帧拼图（每条约 5~20 帧）识别，非原始视频帧；'
        '原始视频全程不落盘。',
        '分析结论由大模型生成并按「推荐度 = 0.40×置信 + 0.25×可操作性 + 0.20×证据充分度 + '
        '0.15×指标一致性」排序，每条结论均绑定可核验的内部证据编号，仅供辅助参考，不输出参考文献。',
        '评论样本未过 30 条闸门的内容，其评论侧结论仅作现象提示。',
        '未覆盖「未知流域」与「黄河长江双命中」样本（另计 %s 条），以保证 长江+黄河 不重复计数。'
        % (meta.get('excluded') or {}),
    ]:
        _p(doc, '· %s' % s, 9.5)

    os.makedirs(REPORTS, exist_ok=True)
    fn = '短视频智能分析简报_%s_%s.docx' % (basin, time.strftime('%Y%m%d_%H%M%S'))
    path = os.path.join(REPORTS, fn)
    doc.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--basin', default='yz', choices=['yz', 'hh'])
    ap.add_argument('--top', type=int, default=10)
    a = ap.parse_args()
    p = build(PFX2BASIN[a.basin], top=a.top)
    print('[report] 已生成：%s  (%.0f KB)' % (p, os.path.getsize(p) / 1024))


if __name__ == '__main__':
    main()
